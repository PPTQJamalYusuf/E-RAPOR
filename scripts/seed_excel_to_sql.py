import openpyxl
import os

excel_path = "MASTER R.TAHFIDZ NEW.xlsx"
output_sql = "supabase/seed.sql"

wb = openpyxl.load_workbook(excel_path, data_only=True)
sheet = wb['Master']

santri_rows = []
for r in range(3, sheet.max_row + 1):
    no = sheet.cell(r, 19).value
    nama = sheet.cell(r, 20).value
    kelas = sheet.cell(r, 21).value
    smt = sheet.cell(r, 22).value
    if nama and str(nama).strip():
        santri_rows.append({
            "no": int(no) if no and str(no).isdigit() else len(santri_rows) + 1,
            "nama": str(nama).strip().replace("'", "''"),
            "kelas": str(kelas).strip().replace("'", "''") if kelas else "الأول",
            "semester": str(smt).strip().replace("'", "''") if smt else "الثاني"
        })

sql_lines = [
    "-- ==============================================================================",
    "-- Seed Data E-Rapot Tahfidz dari 'MASTER R.TAHFIDZ NEW.xlsx'",
    f"-- Total Santri: {len(santri_rows)}",
    "-- ==============================================================================",
    "",
    "-- 1. Tahun Ajaran Default",
    "INSERT INTO tahun_ajaran (id, nama, semester, is_active)",
    "VALUES ('a0000000-0000-0000-0000-000000000001', '2024/2025', 'الثاني', true)",
    "ON CONFLICT DO NOTHING;",
    "",
    "-- 2. Data Santri Master",
]

for s in santri_rows:
    nis = f"ST-{s['no']:04d}"
    sql_lines.append(
        f"INSERT INTO santri (nis, nama, kelas, jenis_kelamin, status) "
        f"VALUES ('{nis}', '{s['nama']}', '{s['kelas']}', 'P', 'aktif') "
        f"ON CONFLICT (nis) DO UPDATE SET nama = EXCLUDED.nama, kelas = EXCLUDED.kelas;"
    )

os.makedirs("supabase", exist_ok=True)
with open(output_sql, "w", encoding="utf-8") as f:
    f.write("\n".join(sql_lines) + "\n")

print(f"Sukses generate {output_sql} dengan {len(santri_rows)} santri.")
